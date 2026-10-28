import csv
import io
import re
from pathlib import Path

from openpyxl import Workbook, load_workbook

MAX_ROWS = 100_000
MAX_BYTES = 50 * 1024 * 1024
SUPPORTED = {".csv", ".xlsx"}
_FORMULA_STARTS = ("=", "+", "-", "@", "\t", "\r")
_PHONE_LIKE = re.compile(r"^[+-]?[\d\s().-]+$")


class TableError(Exception):
    """Raised for input the tool refuses to process."""


def _check_suffix(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise TableError(f"Unsupported file type '{path.suffix}'. Use .csv or .xlsx.")
    return suffix


def _read_csv(path: Path) -> list[list[str]]:
    try:
        raw = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as error:
        raise TableError("The file is not valid UTF-8. Save it as UTF-8 CSV.") from error
    try:
        dialect = csv.Sniffer().sniff(raw[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    try:
        return list(csv.reader(io.StringIO(raw, newline=""), dialect))
    except csv.Error as error:
        raise TableError(f"The CSV could not be parsed: {error}") from error


def _read_xlsx(path: Path) -> list[list[str]]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            records = []
            for row in workbook.worksheets[0].iter_rows(values_only=True):
                records.append(["" if v is None else str(v) for v in row])
                if len(records) > MAX_ROWS + 1:
                    raise TableError(f"Too many rows; the limit is {MAX_ROWS}.")
            return records
        finally:
            workbook.close()
    except TableError:
        raise
    except Exception as error:  # damaged or password-protected workbooks raise many types
        raise TableError(f"The Excel file could not be read: {type(error).__name__}.") from error


def read_table(path: Path) -> tuple[list[str], list[list[str]]]:
    suffix = _check_suffix(path)
    if not path.is_file():
        raise TableError(f"File not found: {path}")
    if path.stat().st_size > MAX_BYTES:
        raise TableError(f"The file is larger than {MAX_BYTES // (1024 * 1024)} MB.")
    records = _read_csv(path) if suffix == ".csv" else _read_xlsx(path)
    records = [row for row in records if any(cell.strip() for cell in row)]
    if not records:
        raise TableError("The file is empty.")
    header, rows = [h.strip() for h in records[0]], records[1:]
    if len(rows) > MAX_ROWS:
        raise TableError(f"Too many rows ({len(rows)}); the limit is {MAX_ROWS}.")
    width = len(header)
    for number, row in enumerate(rows, start=1):
        if any(cell.strip() for cell in row[width:]):
            raise TableError(
                f"Data row {number} has {len(row)} fields but the header has {width}. "
                "Fix the file (an unquoted comma in a value is a common cause); "
                "nothing was changed."
            )
    return header, [(row + [""] * width)[:width] for row in rows]


def _is_number(text: str) -> bool:
    try:
        float(text)
    except ValueError:
        return False
    return True


def needs_neutralising(value: str) -> bool:
    """True for text a spreadsheet program could run as a formula (CSV/formula injection)."""
    if not value.startswith(_FORMULA_STARTS):
        return False
    return not (_is_number(value) or _PHONE_LIKE.match(value))


def sanitize_cell(value: str) -> str:
    return "'" + value if needs_neutralising(value) else value


def write_table(path: Path, header: list[str], rows: list[list[str]]) -> int:
    """Write the table and return how many cells had to be neutralised."""
    suffix = _check_suffix(path)
    cells = [header, *rows]
    neutralised = sum(needs_neutralising(cell) for row in cells for cell in row)
    if suffix == ".csv":
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.writer(handle)
            for row in cells:
                writer.writerow([sanitize_cell(cell) for cell in row])
    else:
        workbook = Workbook()
        sheet = workbook.active
        for row_number, row in enumerate(cells, start=1):
            for column_number, value in enumerate(row, start=1):
                cell = sheet.cell(row=row_number, column=column_number, value=value)
                if needs_neutralising(value):
                    cell.data_type = "s"  # stored as plain text, never as a formula
        workbook.save(path)
    return neutralised
